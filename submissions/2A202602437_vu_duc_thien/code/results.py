"""results.py - Bước 5: gom kết quả các lần chạy vào results.xlsx (GUIDE.md mục 6.1).

Mọi số lấy từ file do chính các lần chạy ghi ra (runs/<exp_id>/seed<k>/result.json) và từ eval.py
(load_group trên predictions/), nên truy ngược được tới exp_id (RUBRIC P2).
"""
from __future__ import annotations

import json
from pathlib import Path

import numpy as np
import pandas as pd

import dataset as D
from eval import load_group  # noqa: E402  (gốc repo đã có trong sys.path khi import train/notebook)

CHINEE, SNAKE = 0, 7


def load_results(out_dir: str = "runs") -> pd.DataFrame:
    """Một dòng cho mỗi lần chạy đã xong (đọc mọi runs/*/seed*/result.json)."""
    rows = []
    for f in sorted(Path(out_dir).glob("*/seed*/result.json")):
        r = json.loads(f.read_text(encoding="utf-8"))
        f1, rec = r.pop("val_f1_per_class"), r.pop("val_recall_per_class")
        r["gpu"] = (r.pop("env", None) or {}).get("gpu")
        r.update({"val_f1_chinee": f1[CHINEE], "val_f1_snake": f1[SNAKE],
                  "val_recall_chinee": rec[CHINEE], "val_recall_snake": rec[SNAKE]})
        rows.append(r)
    return pd.DataFrame(rows)


def backbones_sheet(res: pd.DataFrame) -> pd.DataFrame:
    b = res[res["exp_id"].str.startswith("B")].sort_values("exp_id")
    return pd.DataFrame({
        "exp_id": b["exp_id"], "backbone": b["backbone"].str.split(".").str[0], "tag trọng số": b["weight_tag"],
        "#tham số (M)": b["params_m"], "GMAC": b["gmac"], "độ phân giải": b["img_size"], "epoch": b["epochs"],
        "seed": b["seed"], "macro-F1 val": b["val_macro_f1"], "top-1 val": b["val_top1"],
        "F1 Chinee val": b["val_f1_chinee"], "F1 Snake val": b["val_f1_snake"], "best epoch": b["best_epoch"],
        "thời gian train/epoch (s)": b["train_s_per_epoch"], "độ trễ batch-1 fp32 p50 (ms)": b["latency_b1_fp32_p50_ms"],
        "GPU": b["gpu"], "ghi chú": "công thức nền T00, 1 seed",
    })


def plot_backbones(bs: pd.DataFrame, path) -> Path:
    """macro-F1 val theo độ trễ batch 1 và theo GMAC cho các backbone (Bước 1)."""
    from matplotlib.figure import Figure

    fig = Figure(figsize=(12, 4.5), dpi=120)
    for ax, xcol, xlab in zip(fig.subplots(1, 2), ("độ trễ batch-1 fp32 p50 (ms)", "GMAC"),
                              ("độ trễ p50 batch 1, fp32 (ms)", "GMAC / ảnh 224x224")):
        ax.scatter(bs[xcol], bs["macro-F1 val"], s=20 + 3 * bs["#tham số (M)"])
        for _, r in bs.iterrows():
            ax.annotate(f"{r['exp_id']} {r['backbone']}", (r[xcol], r["macro-F1 val"]), fontsize=7,
                        xytext=(3, 3), textcoords="offset points")
        ax.set(xlabel=xlab, ylabel="macro-F1 val", title=f"Backbone: macro-F1 val theo {xcol.split(' (')[0]}")
        ax.grid(alpha=0.3)
    fig.suptitle("Bước 1 (cùng công thức T00, seed 0; kích thước điểm ~ số tham số)")
    fig.tight_layout()
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, bbox_inches="tight")
    return path


def noise_std(res: pd.DataFrame, exp_id: str = "T00") -> float:
    """std mẫu (ddof=1) macro-F1 val qua các seed của một cấu hình: thước đo nhiễu do seed."""
    v = res.loc[res["exp_id"] == exp_id, "val_macro_f1"]
    return float(v.std(ddof=1)) if len(v) > 1 else float("nan")


def training_sheet(res: pd.DataFrame, design: dict, base: str = "T00") -> pd.DataFrame:
    """Ablation Bước 2. `design`: exp_id -> (trục A-G, khác T00 ở điểm nào). Δ so với T00 cùng seed;
    so với nhiễu = std macro-F1 val của T00 qua các seed."""
    t = res[res["exp_id"].str.startswith("T")].copy()
    ref = t[t["exp_id"] == base].set_index("seed")["val_macro_f1"]
    std = noise_std(res, base)
    t["Δ vs T00"] = t.apply(lambda r: r["val_macro_f1"] - ref.get(r["seed"], ref.iloc[0]), axis=1)
    rows = []
    for _, r in t.sort_values(["exp_id", "seed"]).iterrows():
        axis, diff = design.get(r["exp_id"], ("", ""))
        verdict = "mốc" if r["exp_id"] == base else (
            "trong nhiễu (không phân biệt được)" if not abs(r["Δ vs T00"]) > std
            else ("tốt hơn, vượt nhiễu" if r["Δ vs T00"] > 0 else "kém hơn, vượt nhiễu"))
        note = "mốc; 3 seed để đo nhiễu" if r["exp_id"] == base else "1 seed"
        if r["gpu"] and "T4" not in str(r["gpu"]):
            note += f"; chạy trên {r['gpu']} (các lần khác trên T4)"
        rows.append({"exp_id": r["exp_id"], "backbone": r["backbone"].split(".")[0], "trục": axis,
                     "khác T00 ở điểm nào": diff, "seed": r["seed"], "macro-F1 val": r["val_macro_f1"],
                     "top-1 val": r["val_top1"], "Δ vs T00 (macro-F1)": r["Δ vs T00"],
                     "std T00 qua seed": std, "kết luận so với nhiễu": verdict,
                     "F1 Chinee val": r["val_f1_chinee"], "F1 Snake val": r["val_f1_snake"],
                     "best epoch": r["best_epoch"], "thời gian train/epoch (s)": r["train_s_per_epoch"],
                     "GPU": r["gpu"], "ghi chú": note})
    return pd.DataFrame(rows)


def inference_sheet(inf: pd.DataFrame) -> pd.DataFrame:
    cols = {"exp_id": "exp_id", "method_label": "phương pháp", "model": "mô hình/checkpoint", "K": "K",
            "macro_f1": "macro-F1 val", "top1": "top-1 val", "ece": "ECE val", "f1_chinee": "F1 Chinee val",
            "f1_snake": "F1 Snake val", "lat_b1_p50_ms": "p50 b1 (ms)", "lat_b1_p95_ms": "p95 b1 (ms)",
            "lat_b1_p99_ms": "p99 b1 (ms)", "throughput_b1": "thông lượng b1 (ảnh/s)",
            "rel_cost_vs_I00": "chi phí so với I00 (lần)",
            "temperature": "T (khớp trên val)", "ece_before": "ECE val trước TS",
            "ece_crossfit_after": "ECE val sau TS (khớp chéo 2 nửa)"}
    inf = inf.drop_duplicates(subset=["exp_id", "model"]).copy()  # ensemble/EMA giống nhau giữa các lượt Bước 3
    inf["throughput_b1"] = 1000.0 / inf["lat_b1_p50_ms"]  # batch 1: một ảnh mỗi lượt suy luận (mọi view)
    return inf[[c for c in cols if c in inf.columns]].rename(columns=cols)


def latency_sheet(lat: pd.DataFrame) -> pd.DataFrame:
    """Bảng Latency với tên cột có đơn vị (GUIDE mục 6.1)."""
    cols = {"config": "cấu hình", "gpu": "GPU", "dtype": "dtype", "batch": "batch", "img_size": "độ phân giải",
            "fused_bn": "gộp BN", "k_views": "K (view)", "p50": "p50 (ms)", "p95": "p95 (ms)", "p99": "p99 (ms)",
            "mean": "trung bình (ms)", "images_per_s": "thông lượng (ảnh/s)", "n_iters": "số lần đo",
            "warmup": "warmup (lần)", "torch": "torch", "preprocessing": "tiền xử lý"}
    return lat[[c for c in cols if c in lat.columns]].rename(columns=cols)


def final_group(pattern: str, test_csv: str):
    """eval.load_group: đọc file dự đoán (kiểm tra định dạng, đối chiếu test CSV), chỉ số mỗi seed."""
    return load_group(pattern, test_csv, ref_what="test")


def final_sheet(groups: dict, test_csv: str, val_f1: dict | None = None) -> pd.DataFrame:
    """`groups`: exp_id -> (mô tả cấu hình, glob file test). `val_f1`: exp_id -> {seed: macro-F1 val}."""
    rows = []
    for exp, (desc, pattern) in groups.items():
        g = final_group(pattern, test_csv)
        for p, m in zip(g.preds, g.metrics):
            rows.append({"exp_id": exp, "cấu hình": desc, "seed": p.seed,
                         "macro-F1 val": (val_f1 or {}).get(exp, {}).get(p.seed, np.nan),
                         "macro-F1 test": m["macro_f1"], "top-1 test": m["top1"], "ECE test": m["ece"],
                         "recall Chinee test": m["recall"][CHINEE], "recall Snake test": m["recall"][SNAKE]})
        sub = pd.DataFrame(rows[-len(g.preds):])
        num = sub.select_dtypes("number").drop(columns="seed")

        def agg(v):
            v = v.dropna()
            return "" if v.empty else (f"{v.mean():.4f}" if len(v) == 1 else f"{v.mean():.4f} ± {v.std(ddof=1):.4f}")

        rows.append({"exp_id": exp, "cấu hình": desc, "seed": f"mean ± std ({len(sub)} seed)",
                     **{c: agg(num[c]) for c in num.columns}})
    return pd.DataFrame(rows)


def perclass_sheet(groups: dict, test_csv: str) -> pd.DataFrame:
    """Precision/recall/F1 từng lớp trên test (mean ± std qua seed) cho các nhóm trong `groups`."""
    out = None
    for exp, (_, pattern) in groups.items():
        g = final_group(pattern, test_csv)
        d = pd.DataFrame({"lớp": D.CLASS_NAMES, "số ảnh test": g.metrics[0]["support"].astype(int)})
        for key in ("precision", "recall", "f1"):
            mean, std = g.summary[key]
            d[f"{exp} {key}"] = [f"{a:.4f} ± {b:.4f}" for a, b in zip(mean, std)]
        out = d if out is None else out.merge(d, on=["lớp", "số ảnh test"])
    return out


def summary_sheet(res: pd.DataFrame, inf: pd.DataFrame | None, n: int = 10) -> pd.DataFrame:
    """Top cấu hình theo macro-F1 val (lần chạy huấn luyện và cách suy luận), kèm chi phí/độ trễ."""
    rows = [{"exp_id": f"{r.exp_id}/seed{r.seed}", "loại": "huấn luyện", "mô tả": r.backbone.split(".")[0],
             "macro-F1 val": r.val_macro_f1, "top-1 val": r.val_top1,
             "độ trễ b1 p50 (ms)": r.latency_b1_fp32_p50_ms, "GMAC": r.gmac}
            for r in res.itertuples()]
    if inf is not None:
        rows += [{"exp_id": r.exp_id, "loại": "suy luận", "mô tả": f"{r.method_label} ({r.model})",
                  "macro-F1 val": r.macro_f1, "top-1 val": r.top1, "độ trễ b1 p50 (ms)": r.lat_b1_p50_ms,
                  "GMAC": np.nan} for r in inf.itertuples()]
    s = (pd.DataFrame(rows).drop_duplicates(subset=["exp_id", "mô tả"])
         .sort_values("macro-F1 val", ascending=False).head(n).reset_index(drop=True))
    s.insert(0, "hạng", np.arange(1, len(s) + 1))
    return s


def write_xlsx(path, sheets: dict, notes: dict | None = None) -> Path:
    """Ghi các sheet, cố định hàng tiêu đề, 4 chữ số thập phân, tô đậm dòng có macro-F1 cao nhất."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    best_fill = PatternFill("solid", fgColor="FFF2CC")
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            start = 2 if notes and name in notes else 0
            df.to_excel(xw, sheet_name=name, index=False, startrow=start)
            ws = xw.sheets[name]
            if start:
                ws.cell(row=1, column=1, value=notes[name]).font = Font(italic=True)
            ws.freeze_panes = ws.cell(row=start + 2, column=1)
            for j, col in enumerate(df.columns, 1):
                ws.cell(row=start + 1, column=j).font = Font(bold=True)
                width = max([len(str(col))] + [len(f"{v:.4f}" if isinstance(v, float) else str(v))
                                               for v in df[col].head(200)])
                ws.column_dimensions[get_column_letter(j)].width = min(48, width + 2)
                if pd.api.types.is_float_dtype(df[col]):
                    for i in range(len(df)):
                        ws.cell(row=start + 2 + i, column=j).number_format = "0.0000"
            key = next((c for c in df.columns if str(c).startswith("macro-F1")
                        and pd.api.types.is_numeric_dtype(df[c])), None)
            if key is not None and df[key].notna().any():
                best = int(df[key].astype(float).idxmax())
                for j in range(1, len(df.columns) + 1):
                    ws.cell(row=start + 2 + df.index.get_loc(best), column=j).fill = best_fill
    return path
